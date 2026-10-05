"""Da formato de petición a una hoja Excel con la estructura de EJEMPLO INICIAL.

Requiere openpyxl: python -m pip install openpyxl

Uso:
    python formatear_peticion.py entrada.xlsx
    python formatear_peticion.py entrada.xlsx -o salida.xlsx
"""

from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path
import sys
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator, TranslatorError
from openpyxl.styles import Border
from openpyxl.utils.exceptions import InvalidFileException


HOJA = "PETICION OFERTAS"
NOMBRE_PLANTILLA = "EJEMPLO PETICIÓN ELECTRICIDAD.xlsx"
PRIMERA_FILA_DATOS_ORIGEN = 7
PRIMERA_FILA_DATOS_SALIDA = 8
PRIMERA_COLUMNA_DATOS = 2
ULTIMA_COLUMNA_DATOS = 33


def carpeta_aplicacion() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def ruta_recurso(nombre: str) -> Path:
    carpeta = Path(getattr(sys, "_MEIPASS", carpeta_aplicacion()))
    return carpeta / nombre


def copiar_fila(origen, destino, fila_origen: int, fila_destino: int) -> None:
    for columna in range(PRIMERA_COLUMNA_DATOS, ULTIMA_COLUMNA_DATOS + 1):
        celda_origen = origen.cell(fila_origen, columna)
        celda_destino = destino.cell(fila_destino, columna)
        valor = celda_origen.value

        if isinstance(valor, str) and valor.startswith("="):
            try:
                valor = Translator(
                    valor, origin=celda_origen.coordinate
                ).translate_formula(celda_destino.coordinate)
            except TranslatorError as error:
                raise ValueError(
                    f"No se pudo adaptar la fórmula de {celda_origen.coordinate}: "
                    f"{valor}"
                ) from error

        celda_destino.value = valor
        if celda_origen.hyperlink:
            celda_destino._hyperlink = copy(celda_origen.hyperlink)
        if celda_origen.comment:
            celda_destino.comment = copy(celda_origen.comment)


def formatear(
    entrada: Path, plantilla: Path, salida: Path, *, sobrescribir: bool = False
) -> None:
    entrada_resuelta = entrada.resolve()
    plantilla_resuelta = plantilla.resolve()
    salida_resuelta = salida.resolve()
    if salida_resuelta in (entrada_resuelta, plantilla_resuelta):
        raise ValueError("La salida no puede sobrescribir la entrada ni la plantilla.")
    if entrada_resuelta == plantilla_resuelta:
        raise ValueError("La entrada y la plantilla deben ser archivos distintos.")
    if salida.suffix.lower() != ".xlsx":
        raise ValueError("El archivo de salida debe tener extensión .xlsx.")
    if not entrada.is_file():
        raise FileNotFoundError(f"No existe el archivo de entrada: {entrada}")
    if not plantilla.is_file():
        raise FileNotFoundError(f"No existe la plantilla: {plantilla}")
    if salida.exists() and not sobrescribir:
        raise FileExistsError(
            f"Ya existe el archivo de salida: {salida}. "
            "Indica otra ruta para evitar sobrescribirlo."
        )

    libro_origen = load_workbook(entrada, data_only=False)
    libro_salida = load_workbook(plantilla, data_only=False)
    if HOJA not in libro_origen.sheetnames or HOJA not in libro_salida.sheetnames:
        raise ValueError(f"Ambos archivos deben contener la hoja «{HOJA}».")
    if len(libro_origen.worksheets) != 1:
        raise ValueError("El archivo de entrada debe contener una sola hoja.")

    hoja_origen = libro_origen[HOJA]
    hoja_salida = libro_salida[HOJA]

    filas_con_datos = [
        fila
        for fila in range(PRIMERA_FILA_DATOS_ORIGEN, hoja_origen.max_row + 1)
        if any(
            hoja_origen.cell(fila, columna).value is not None
            for columna in range(PRIMERA_COLUMNA_DATOS, ULTIMA_COLUMNA_DATOS + 1)
        )
    ]
    if not filas_con_datos:
        raise ValueError("No se encontraron datos a partir de la fila 7.")

    grupos = []
    clave_anterior = None
    for indice, fila in enumerate(filas_con_datos):
        clave = tuple(hoja_origen.cell(fila, columna).value for columna in (2, 3, 4))
        inicio_grupo = clave != clave_anterior
        siguiente_clave = (
            tuple(
                hoja_origen.cell(filas_con_datos[indice + 1], columna).value
                for columna in (2, 3, 4)
            )
            if indice + 1 < len(filas_con_datos)
            else None
        )
        fin_grupo = siguiente_clave != clave
        grupos.append((fila, inicio_grupo, fin_grupo))
        clave_anterior = clave

    for fila in range(PRIMERA_FILA_DATOS_ORIGEN, hoja_origen.max_row + 1):
        for columna in (1, *range(ULTIMA_COLUMNA_DATOS + 1, hoja_origen.max_column + 1)):
            if hoja_origen.cell(fila, columna).value is not None:
                raise ValueError(
                    "Se encontraron datos fuera de las columnas B:AG "
                    f"(celda {hoja_origen.cell(fila, columna).coordinate})."
                )

    filas_salida = []
    fila_destino = PRIMERA_FILA_DATOS_SALIDA
    for fila_origen, inicio_grupo, fin_grupo in grupos:
        if inicio_grupo and filas_salida:
            filas_salida.append((None, fila_destino, False, False))
            fila_destino += 1
        filas_salida.append((fila_origen, fila_destino, inicio_grupo, fin_grupo))
        fila_destino += 1
    ultima_fila_salida = fila_destino - 1
    estilo_primera_fila = [
        copy(hoja_salida.cell(PRIMERA_FILA_DATOS_SALIDA, col)._style)
        for col in range(1, ULTIMA_COLUMNA_DATOS + 1)
    ]
    estilo_fila_datos = [
        copy(hoja_salida.cell(PRIMERA_FILA_DATOS_SALIDA + 1, col)._style)
        for col in range(1, ULTIMA_COLUMNA_DATOS + 1)
    ]
    estilo_fin_grupo = [
        copy(hoja_salida.cell(17, col)._style)
        for col in range(1, ULTIMA_COLUMNA_DATOS + 1)
    ]
    borde_izquierdo_exterior = copy(hoja_salida["B8"].border.left)
    borde_derecho_exterior = copy(hoja_salida["P8"].border.right)
    borde_superior_exterior = copy(hoja_salida["B8"].border.top)
    borde_inferior_primera_fila = copy(hoja_salida["B8"].border.bottom)
    borde_inferior_exterior = copy(hoja_salida["B17"].border.bottom)
    alto_primera_fila = hoja_salida.row_dimensions[
        PRIMERA_FILA_DATOS_SALIDA
    ].height
    alto_fila_datos = hoja_salida.row_dimensions[
        PRIMERA_FILA_DATOS_SALIDA + 1
    ].height
    estilo_total = [
        copy(hoja_salida.cell(208, col)._style)
        for col in range(1, ULTIMA_COLUMNA_DATOS + 1)
    ]
    estilo_total_segunda_fila = [
        copy(hoja_salida.cell(209, col)._style)
        for col in range(1, ULTIMA_COLUMNA_DATOS + 1)
    ]
    alto_total = hoja_salida.row_dimensions[208].height
    alto_total_segunda_fila = hoja_salida.row_dimensions[209].height

    for rango in list(hoja_salida.merged_cells.ranges):
        if rango.min_row >= PRIMERA_FILA_DATOS_SALIDA:
            hoja_salida.unmerge_cells(str(rango))

    for fila in range(PRIMERA_FILA_DATOS_SALIDA, hoja_salida.max_row + 1):
        for columna in range(1, hoja_salida.max_column + 1):
            celda = hoja_salida.cell(fila, columna)
            celda.value = None
            celda.hyperlink = None
            celda.comment = None

    for fila_origen, fila_destino, inicio_grupo, fin_grupo in filas_salida:
        if fila_origen is not None:
            copiar_fila(hoja_origen, hoja_salida, fila_origen, fila_destino)
            if not inicio_grupo:
                for columna in (*range(2, 17), 24, 25):
                    celda = hoja_salida.cell(fila_destino, columna)
                    celda.value = None
                    celda.hyperlink = None
                    celda.comment = None
            else:
                siguiente_separador = next(
                    (
                        fila
                        for origen, fila, _, _ in filas_salida
                        if origen is None and fila > fila_destino
                    ),
                    ultima_fila_salida + 1,
                )
                hoja_salida.cell(fila_destino, 24).value = (
                    f"=SUM(R{fila_destino}:W{siguiente_separador - 1})"
                )

        estilos = estilo_primera_fila if inicio_grupo else estilo_fila_datos
        if fin_grupo and not inicio_grupo:
            estilos = [
                fin_style if columna <= 23 else style
                for columna, (style, fin_style) in enumerate(
                    zip(estilos, estilo_fin_grupo), start=1
                )
            ]
        for columna, estilo in enumerate(estilos, start=1):
            hoja_salida.cell(fila_destino, columna)._style = copy(estilo)
        if fila_origen is None:
            for columna in range(1, ULTIMA_COLUMNA_DATOS + 1):
                hoja_salida.cell(fila_destino, columna).border = Border()
        elif not inicio_grupo:
            for columna in range(24, ULTIMA_COLUMNA_DATOS + 1):
                hoja_salida.cell(fila_destino, columna).border = Border()
        if fila_origen is not None:
            for columna in range(24, ULTIMA_COLUMNA_DATOS + 1):
                celda = hoja_salida.cell(fila_destino, columna)
                borde_estilo = celda.border
                celda.border = Border(
                    left=(
                        copy(borde_izquierdo_exterior)
                        if columna == 24
                        else copy(borde_estilo.left) if inicio_grupo else None
                    ),
                    right=(
                        copy(borde_derecho_exterior)
                        if columna == ULTIMA_COLUMNA_DATOS
                        else copy(borde_estilo.right) if inicio_grupo else None
                    ),
                    top=copy(borde_superior_exterior) if inicio_grupo else None,
                    bottom=(
                        copy(borde_inferior_primera_fila)
                        if inicio_grupo
                        else copy(borde_inferior_exterior) if fin_grupo else None
                    ),
                )
        alto = alto_primera_fila if inicio_grupo else alto_fila_datos
        if alto is not None:
            hoja_salida.row_dimensions[fila_destino].height = alto

    if hoja_salida.max_row > ultima_fila_salida:
        hoja_salida.delete_rows(
            ultima_fila_salida + 1, hoja_salida.max_row - ultima_fila_salida
        )
    for fila in list(hoja_salida.row_dimensions):
        if fila > ultima_fila_salida:
            del hoja_salida.row_dimensions[fila]

    total_fila = ultima_fila_salida + 4
    for fila in range(ultima_fila_salida + 1, total_fila + 2):
        for columna in range(1, ULTIMA_COLUMNA_DATOS + 1):
            celda = hoja_salida.cell(fila, columna)
            celda.value = None
            celda._style = None
            celda.hyperlink = None
            celda.comment = None
    for columna, estilo in enumerate(estilo_total, start=1):
        hoja_salida.cell(total_fila, columna)._style = copy(estilo)
    for columna, estilo in enumerate(estilo_total_segunda_fila, start=1):
        hoja_salida.cell(total_fila + 1, columna)._style = copy(estilo)
    hoja_salida.cell(total_fila, 17).value = "TOTAL"
    hoja_salida.cell(total_fila, 23).value = "=SUM(X:X)"
    if alto_total is not None:
        hoja_salida.row_dimensions[total_fila].height = alto_total
    if alto_total_segunda_fila is not None:
        hoja_salida.row_dimensions[total_fila + 1].height = alto_total_segunda_fila
    hoja_salida.merge_cells(
        start_row=total_fila, start_column=17, end_row=total_fila + 1, end_column=22
    )
    hoja_salida.merge_cells(
        start_row=total_fila, start_column=23, end_row=total_fila + 1, end_column=23
    )

    salida.parent.mkdir(parents=True, exist_ok=True)
    libro_salida.save(salida)


def ejecutar_linea_comandos() -> None:
    parser = argparse.ArgumentParser(
        description="Aplica el formato de petición a un archivo Excel."
    )
    parser.add_argument("entrada", type=Path, help="Archivo Excel con los datos.")
    parser.add_argument(
        "-o",
        "--salida",
        type=Path,
        help="Ruta del archivo formateado (por defecto, añade _peticion al nombre).",
    )
    parser.add_argument(
        "-t",
        "--plantilla",
        type=Path,
        default=Path(__file__).with_name(NOMBRE_PLANTILLA),
        help="Plantilla de formato (por defecto, el ejemplo de petición de esta carpeta).",
    )
    argumentos = parser.parse_args()
    salida = argumentos.salida or argumentos.entrada.with_name(
        f"{argumentos.entrada.stem}_peticion.xlsx"
    )

    try:
        formatear(argumentos.entrada, argumentos.plantilla, salida)
    except (FileExistsError, FileNotFoundError, ValueError, OSError) as error:
        parser.error(str(error))
    print(f"Archivo creado: {salida}")


def ejecutar_interfaz() -> None:
    ventana = tk.Tk()
    ventana.title("Formato de petición")
    ventana.resizable(False, False)
    ventana.minsize(620, 210)

    archivo_entrada = tk.StringVar()
    carpeta_salida = tk.StringVar(value=str(carpeta_aplicacion()))
    estado = tk.StringVar(value="Selecciona el archivo Excel que quieres formatear.")

    marco = ttk.Frame(ventana, padding=18)
    marco.grid(sticky="nsew")
    marco.columnconfigure(1, weight=1)

    ttk.Label(marco, text="Archivo Excel de entrada:").grid(
        row=0, column=0, sticky="w", padx=(0, 10), pady=(0, 8)
    )
    ttk.Entry(marco, textvariable=archivo_entrada, width=56).grid(
        row=0, column=1, sticky="ew", pady=(0, 8)
    )

    def elegir_entrada() -> None:
        elegido = filedialog.askopenfilename(
            parent=ventana,
            title="Seleccionar archivo Excel",
            filetypes=[("Archivos Excel", "*.xlsx"), ("Todos los archivos", "*.*")],
        )
        if elegido:
            archivo_entrada.set(elegido)

    ttk.Button(marco, text="Examinar...", command=elegir_entrada).grid(
        row=0, column=2, padx=(8, 0), pady=(0, 8)
    )

    ttk.Label(marco, text="Carpeta del resultado:").grid(
        row=1, column=0, sticky="w", padx=(0, 10), pady=8
    )
    ttk.Entry(marco, textvariable=carpeta_salida, width=56).grid(
        row=1, column=1, sticky="ew", pady=8
    )

    def elegir_carpeta() -> None:
        elegido = filedialog.askdirectory(
            parent=ventana,
            title="Elegir carpeta del resultado",
            initialdir=carpeta_salida.get() or str(carpeta_aplicacion()),
        )
        if elegido:
            carpeta_salida.set(elegido)

    ttk.Button(marco, text="Examinar...", command=elegir_carpeta).grid(
        row=1, column=2, padx=(8, 0), pady=8
    )
    ttk.Label(
        marco, text="El resultado se llamará <archivo>_peticion.xlsx."
    ).grid(row=2, column=0, columnspan=3, sticky="w", pady=(0, 12))

    def iniciar_formateo() -> None:
        entrada = Path(archivo_entrada.get().strip())
        directorio = Path(carpeta_salida.get().strip())
        if not entrada.is_file():
            messagebox.showerror(
                "Archivo no válido",
                "Selecciona un archivo Excel de entrada válido.",
                parent=ventana,
            )
            return
        if not directorio.is_dir():
            messagebox.showerror(
                "Carpeta no válida",
                "Selecciona una carpeta de resultado válida.",
                parent=ventana,
            )
            return

        salida = directorio / f"{entrada.stem}_peticion.xlsx"
        sobrescribir = False
        if salida.exists():
            sobrescribir = messagebox.askyesno(
                "El archivo ya existe",
                f"¿Quieres reemplazar este archivo?\n\n{salida}",
                parent=ventana,
            )
            if not sobrescribir:
                return

        boton_formatear.configure(state="disabled")
        estado.set("Aplicando formato...")
        ventana.update_idletasks()
        try:
            formatear(
                entrada,
                ruta_recurso(NOMBRE_PLANTILLA),
                salida,
                sobrescribir=sobrescribir,
            )
        except (
            BadZipFile,
            FileExistsError,
            FileNotFoundError,
            InvalidFileException,
            ValueError,
            OSError,
        ) as error:
            estado.set("No se pudo crear el archivo.")
            messagebox.showerror("Error al formatear", str(error), parent=ventana)
        else:
            estado.set(f"Archivo creado: {salida}")
            messagebox.showinfo(
                "Formato completado",
                f"El archivo se guardó correctamente en:\n\n{salida}",
                parent=ventana,
            )
        finally:
            boton_formatear.configure(state="normal")

    boton_formatear = ttk.Button(
        marco, text="Formatear archivo", command=iniciar_formateo
    )
    boton_formatear.grid(row=3, column=0, columnspan=3, pady=(0, 8))
    ttk.Label(marco, textvariable=estado, wraplength=570).grid(
        row=4, column=0, columnspan=3, sticky="w"
    )
    ventana.mainloop()


if __name__ == "__main__":
    if len(sys.argv) == 1:
        ejecutar_interfaz()
    else:
        ejecutar_linea_comandos()
